# -*- coding: utf-8 -*-
# Проект Морбековой Камиллы. Генератор анализа и материалов защиты.
# Этот код подготовлен в чате, но не был выполнен там из-за сбоя среды.
# Все численные результаты вычисляются при запуске; экспертные оценки не выдумываются.
from __future__ import annotations
import argparse, gc, hashlib, html, importlib.metadata, json, platform, shutil, sys, time, zipfile
from datetime import datetime, timezone
from pathlib import Path
from statistics import median
import numpy as np
import pandas as pd
import psutil
import duckdb
import joblib
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from scipy import sparse
from sklearn.feature_extraction.text import TfidfVectorizer
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.xmlchemy import OxmlElement
from docx import Document
from docx.shared import Pt as DocPt, Cm


AUTHOR = 'Морбекова Камилла'
DIPLOMA = 'Разработка ИИ-агента для преподавателя, исследователя'
DATA_URL = 'https://huggingface.co/datasets/yufan/arxiv-metadata-2020-2026'
SOURCES = [
    ('Набор данных: yufan, arXiv + Semantic Scholar', DATA_URL),
    ('Структура идентификаторов arXiv', 'https://info.arxiv.org/help/arxiv_identifier.html'),
    ('TF-IDF, официальная документация', 'https://scikit-learn.org/stable/modules/generated/sklearn.feature_extraction.text.TfidfVectorizer.html'),
    ('Косинусное сходство', 'https://scikit-learn.org/stable/modules/generated/sklearn.metrics.pairwise.cosine_similarity.html'),
    ('DuckDB: чтение Parquet', 'https://duckdb.org/docs/stable/data/parquet/overview.html'),
    ('Требования курса', 'Занятие 2: слайды 17–19; занятие 3: слайд 7; занятие 4: слайды 8–10; занятие 5: слайды 2, 7, 10–12. Формат заменён на презентацию по сообщению студентки.')
]
QUERIES = [
    ('Q01', 'large language models in education', 'Большие языковые модели в образовании'),
    ('Q02', 'intelligent tutoring systems student feedback', 'Обучающие системы и обратная связь студентам'),
    ('Q03', 'automated essay scoring feedback', 'Автоматическая оценка эссе и обратная связь'),
    ('Q04', 'automatic question generation educational assessment', 'Генерация вопросов для проверки знаний'),
    ('Q05', 'retrieval augmented generation scientific literature', 'Поиск научной литературы для RAG'),
    ('Q06', 'scientific paper recommendation', 'Рекомендация научных публикаций'),
    ('Q07', 'learning analytics student performance prediction', 'Учебная аналитика и успеваемость'),
    ('Q08', 'hallucination detection factual consistency large language models', 'Проверка фактической достоверности ответов'),
    ('Q09', 'academic integrity AI generated text detection', 'Академическая честность и тексты ИИ'),
    ('Q10', 'personalized learning recommendation systems', 'Персонализированное обучение')
]




def write_json(path, data):
    def convert(x):
        if isinstance(x, np.integer): return int(x)
        if isinstance(x, np.floating): return None if not np.isfinite(x) else float(x)
        if isinstance(x, Path): return str(x)
        raise TypeError(str(type(x)))
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2, default=convert), encoding='utf-8')




def show_num(x, digits=0):
    if x is None or (isinstance(x, float) and not np.isfinite(x)): return 'нет данных'
    return f'{x:,.{digits}f}'.replace(',', ' ').replace('.', ',')




def missing(series):
    return series.isna() | series.astype(str).str.strip().eq('')




def save_table(frame, path):
    frame.to_csv(path, index=False, encoding='utf-8-sig')




def benchmark(fun, repeats=3):
    times = []
    for _ in range(repeats):
        gc.collect()
        started = time.perf_counter()
        result = fun()
        times.append(time.perf_counter() - started)
        del result
    return median(times), times




def evaluate_labels(root, summary):
    results = pd.read_csv(root / 'results/search_results.csv', dtype=str, keep_default_na=False)
    path = root / 'evaluation/relevance_labels.csv'
    if not path.exists():
        return {'status': 'pending', 'message': 'Ручная оценка ещё не выполнена.'}
    labels = pd.read_csv(path, dtype=str, keep_default_na=False)
    if 'relevant' in labels: labels['relevant'] = labels['relevant'].str.strip()
    keys = ['query_id', 'arxiv_id']
    required = keys + ['source_sha256', 'relevant', 'reviewer', 'reason']
    if not set(required).issubset(labels.columns):
        raise ValueError('В файле разметки отсутствуют обязательные колонки.')
    if labels.duplicated(keys).any():
        raise ValueError('В разметке повторяется пара запрос–публикация.')
    allowed = {'', '0', '1'}
    if not set(labels.relevant.str.strip()).issubset(allowed):
        raise ValueError('В relevant разрешены только 0, 1 или пустая ячейка.')
    judged = labels.relevant.str.strip().isin(['0', '1'])
    if (labels.loc[judged, 'source_sha256'] != summary['source_sha256']).any():
        raise ValueError('Разметка относится к другому исходному файлу.')
    check = results[keys].merge(labels[required], on=keys, how='left', validate='one_to_one')
    complete = check['relevant'].isin(['0', '1']) & check['reviewer'].fillna('').str.strip().ne('') & check['reason'].fillna('').str.strip().ne('')
    if not len(check) or not complete.all():
        return {'status': 'pending', 'judged': int(complete.sum()), 'returned': len(check), 'message': 'Нет полной ручной проверки: Precision@10 не объявляется.'}
    rows = []
    for qid, query, ru in QUERIES:
        part = check[check.query_id == qid]
        positives = int((part.relevant == '1').sum())
        rows.append({'query_id': qid, 'returned': len(part), 'relevant': positives, 'precision_at_10': positives / 10})
    metrics = pd.DataFrame(rows)
    save_table(metrics, root / 'results/relevance_metrics.csv')
    return {'status': 'complete', 'mean_precision_at_10': float(metrics.precision_at_10.mean()), 'queries_meeting_7_of_10': int((metrics.relevant >= 7).sum()), 'queries': len(QUERIES), 'reviewers': sorted(set(check.reviewer)), 'message': 'Оценка по введённым человеком меткам; квалификация проверяющего отдельно не подтверждалась.'}




def run_analysis(input_path, root):
    for sub in ['data/raw', 'data/processed', 'data/features', 'results/figures', 'evaluation', 'index', 'materials']:
        (root / sub).mkdir(parents=True, exist_ok=True)
    raw_path = root / 'data/raw/metadata.jsonl'
    if input_path.resolve() != raw_path.resolve(): shutil.copyfile(input_path, raw_path)
    digest = hashlib.sha256()
    with raw_path.open('rb') as source:
        for block in iter(lambda: source.read(1024 * 1024), b''): digest.update(block)
    records, empty_lines = [], 0
    started = time.perf_counter()
    with raw_path.open('r', encoding='utf-8-sig') as source:
        for line_no, line in enumerate(source, 1):
            if not line.strip():
                empty_lines += 1
                continue
            try: item = json.loads(line)
            except json.JSONDecodeError as exc: raise ValueError(f'Некорректный JSON в строке {line_no}: {exc}') from exc
            if not isinstance(item, dict): raise ValueError(f'Строка {line_no} — не JSON-объект.')
            records.append(item)
    read_json_seconds = time.perf_counter() - started
    if not records: raise ValueError('Файл не содержит записей.')
    raw = pd.DataFrame(records)
    del records
    required = ['arxiv_id', 'title', 'authors', 'year', 'publicationDate', 'citationCount', 'influentialCitationCount', 'fieldsOfStudy', 'abstract']
    absent = sorted(set(required) - set(raw.columns))
    if absent: raise ValueError(f'Не найдены поля: {absent}')
    n = len(raw)
    quality = pd.DataFrame({'field': raw.columns, 'missing_count': [int(missing(raw[c]).sum()) for c in raw.columns]})
    quality['missing_pct'] = quality.missing_count / n * 100
    save_table(quality, root / 'results/missing_values.csv')
    save_table(pd.DataFrame({'field': raw.columns, 'dtype_observed': [str(raw[c].dtype) for c in raw.columns]}), root / 'results/schema.csv')
    summary = {
        'author': AUTHOR, 'diploma': DIPLOMA,
        'run_utc': datetime.now(timezone.utc).isoformat(),
        'source_url': DATA_URL, 'source_sha256': digest.hexdigest(),
        'raw_rows': n, 'raw_columns': len(raw.columns), 'empty_lines': empty_lines,
        'raw_jsonl_MB': raw_path.stat().st_size / 1e6,
        'json_parse_seconds': read_json_seconds,
        'full_duplicates': int(raw.astype(str).duplicated().sum()),
        'missing_abstract_raw': int(missing(raw.abstract).sum()),
        'source_description': 'Предоставленный файл из Computer_Science/2025; year не принудительно заменяется на 2025.'
    }
    print(f'Загружено {n:,} записей, {len(raw.columns)} полей.', flush=True)
    clean = raw.copy()
    clean['_source_row'] = np.arange(n)
    for col in raw.columns:
        if col not in ['corpusId', 'year', 'citationCount', 'influentialCitationCount']:
            clean[col] = clean[col].astype('string').str.strip()
            clean.loc[clean[col].eq(''), col] = pd.NA
    clean['arxiv_id_original'] = clean.arxiv_id
    clean['arxiv_id'] = clean.arxiv_id.str.replace(r'^arXiv:', '', regex=True).str.replace(r'v\d+$', '', regex=True)
    valid_id = clean.arxiv_id.str.fullmatch(r'\d{2}(?:0[1-9]|1[0-2])\.\d{4,5}', na=False)
    summary['invalid_or_missing_id'] = int((~valid_id).sum())
    summary['missing_title_raw'] = int(missing(raw.title).sum())
    summary['duplicate_id_extra_rows'] = int(clean.loc[valid_id, 'arxiv_id'].duplicated().sum())
    clean['abstract_words'] = clean.abstract.fillna('').str.split().str.len()
    clean = clean.loc[valid_id & ~missing(clean.title)].copy()
    before_dedup = len(clean)
    clean = clean.sort_values(['abstract_words', '_source_row'], ascending=[False, True], kind='stable').drop_duplicates('arxiv_id').sort_values('_source_row').reset_index(drop=True)
    if not len(clean): raise ValueError('После проверки ключа и названия не осталось записей.')
    summary['duplicates_removed'] = before_dedup - len(clean)
    numeric_rows = []
    for col in ['year', 'citationCount', 'influentialCitationCount']:
        values = pd.to_numeric(clean[col], errors='coerce')
        invalid = values.notna() & ((values < 0) | ~np.isfinite(values) | (values % 1 != 0))
        if col == 'year': invalid = invalid | (values.notna() & ((values < 1900) | (values > datetime.now(timezone.utc).year + 1)))
        malformed = ~missing(clean[col]) & values.isna()
        numeric_rows.append({'field': col, 'min': values[~invalid].min(), 'median': values[~invalid].median(), 'max': values[~invalid].max(), 'invalid_count': int(invalid.sum() + malformed.sum())})
        clean[col] = values.mask(invalid).astype('Float64')
    numeric = pd.DataFrame(numeric_rows)
    save_table(numeric, root / 'results/numeric_quality.csv')
    summary['invalid_numeric_count'] = int(numeric.invalid_count.sum())
    dates = pd.to_datetime(clean.publicationDate, errors='coerce', utc=True)
    summary['invalid_dates'] = int((~missing(clean.publicationDate) & dates.isna()).sum())
    summary['min_publication_date'] = str(dates.min().date()) if dates.notna().any() else 'нет'
    summary['max_publication_date'] = str(dates.max().date()) if dates.notna().any() else 'нет'
    summary['year_not_2025'] = int((clean.year.notna() & clean.year.ne(2025)).sum())
    summary['year_date_mismatches'] = int((clean.year.notna() & dates.notna() & clean.year.ne(dates.dt.year)).sum())
    clean['arxiv_month'] = '20' + clean.arxiv_id.str[:2] + '-' + clean.arxiv_id.str[2:4]
    clean['publication_month'] = dates.dt.strftime('%Y-%m')
    clean['has_abstract'] = ~missing(clean.abstract)
    clean['title_words'] = clean.title.fillna('').str.split().str.len()
    clean['author_count'] = clean.authors.fillna('').map(lambda x: len([a for a in x.split(';') if a.strip()]))
    clean['field_count'] = clean.fieldsOfStudy.fillna('').map(lambda x: len({a.strip() for a in x.split(';') if a.strip()}))
    summary['year_counts'] = {str(k): int(v) for k, v in clean.year.value_counts(dropna=False).sort_index().items()}
    summary['title_duplicates_different_ids'] = int(clean.title.str.casefold().str.replace(r'\s+', ' ', regex=True).duplicated().sum())
    summary['influential_exceeds_total'] = int((clean.influentialCitationCount > clean.citationCount).fillna(False).sum())
    summary['clean_rows'] = len(clean)
    summary['rows_removed'] = n - len(clean)
    summary['abstract_missing_clean'] = int((~clean.has_abstract).sum())
    summary['abstract_missing_pct'] = 100 * summary['abstract_missing_clean'] / len(clean)
    summary['citation_median'] = None if clean.citationCount.notna().sum() == 0 else float(clean.citationCount.median())
    summary['citation_max'] = None if clean.citationCount.notna().sum() == 0 else float(clean.citationCount.max())
    summary['zero_citations'] = int(clean.citationCount.eq(0).sum())
    summary['author_median'] = float(clean.author_count.median())
    summary['abstract_words_median'] = float(clean.loc[clean.has_abstract, 'abstract_words'].median()) if clean.has_abstract.any() else None
    summary['minimum_100k_satisfied'] = len(clean) >= 100000
    if not len(clean): raise ValueError('После проверки ключа и названия не осталось записей.')
    summary['month_counts'] = {str(k): int(v) for k, v in clean.arxiv_month.value_counts().sort_index().items()}
    expected_months = {f'2025-{month:02d}' for month in range(1, 13)}
    summary['missing_arxiv_months_2025'] = sorted(expected_months - set(clean.arxiv_month))
    summary['arxiv_months_outside_2025'] = sorted(set(clean.arxiv_month) - expected_months)
    summary['corpus_id_extra_rows'] = int(clean.loc[clean.corpusId.notna(), 'corpusId'].duplicated().sum()) if 'corpusId' in clean else None
    summary['field_values_original'] = {str(k): int(v) for k, v in clean.fieldsOfStudy.value_counts(dropna=False).items()}
    save_table(pd.DataFrame(sorted(summary['month_counts'].items()), columns=['arxiv_month', 'publications']), root / 'results/month_coverage.csv')
    months = clean.groupby('arxiv_month', as_index=False).agg(month_paper_count=('arxiv_id', 'size'), month_mean_citations=('citationCount', 'mean'))
    before_join = len(clean)
    features = clean.merge(months, on='arxiv_month', how='left', validate='many_to_one', sort=False)
    assert len(features) == before_join and features.arxiv_id.is_unique
    summary['join_rows_before'] = before_join
    summary['join_rows_after'] = len(features)
    summary['join_description'] = 'Обогащение агрегатами того же корпуса по arxiv_month; не является вторым независимым источником.'
    parquet_path = root / 'data/processed/publications.parquet'
    csv_path = root / 'data/processed/publications.csv'
    clean.to_parquet(parquet_path, index=False, engine='pyarrow', compression='snappy')
    clean.to_csv(csv_path, index=False)
    features.to_parquet(root / 'data/features/publications_features.parquet', index=False, compression='snappy')
    save_table(months, root / 'results/monthly_summary.csv')
    text_cols = list(clean.select_dtypes(include=['object', 'string']).columns)
    csv_reader = lambda: pd.read_csv(csv_path, dtype={c: 'string' for c in text_cols}, low_memory=False, keep_default_na=False, na_values=[''])
    parquet_reader = lambda: pd.read_parquet(parquet_path)
    loaded_csv, loaded_parquet = csv_reader(), parquet_reader()
    assert len(loaded_csv) == len(loaded_parquet) == len(clean)
    assert loaded_csv.arxiv_id.tolist() == loaded_parquet.arxiv_id.tolist()
    pd.testing.assert_frame_equal(loaded_csv.reset_index(drop=True), loaded_parquet.reset_index(drop=True), check_dtype=False, check_exact=False)
    del loaded_csv, loaded_parquet
    print('Проверка качества завершена. Измеряю форматы и агрегации...', flush=True)
    tcsv, csv_runs = benchmark(csv_reader)
    tpq, pq_runs = benchmark(parquet_reader)
    connection = duckdb.connect()
    connection.execute('SET threads=1')
    sql_path = str(parquet_path.resolve()).replace("'", "''")
    sql = f"SELECT arxiv_month, count(*) AS month_paper_count, avg(citationCount) AS month_mean_citations FROM read_parquet('{sql_path}') GROUP BY arxiv_month ORDER BY arxiv_month"
    td, duck_runs = benchmark(lambda: connection.execute(sql).df())
    tp, pandas_agg_runs = benchmark(lambda: pd.read_parquet(parquet_path, columns=['arxiv_month', 'citationCount']).groupby('arxiv_month', as_index=False).agg(month_paper_count=('citationCount', 'size'), month_mean_citations=('citationCount', 'mean')))
    actual_duck = connection.execute(sql).df()
    pd.testing.assert_frame_equal(months.sort_values('arxiv_month').reset_index(drop=True), actual_duck, check_dtype=False, check_exact=False, rtol=1e-9, atol=1e-9)
    connection.close()
    env = {'os': platform.platform(), 'python': platform.python_version(), 'logical_cpus': psutil.cpu_count(), 'system_RAM_GiB': psutil.virtual_memory().total / 1024 ** 3, 'process_RSS_GiB': psutil.Process().memory_info().rss / 1024 ** 3}
    for candidate in ['/sys/fs/cgroup/memory.max', '/sys/fs/cgroup/memory/memory.limit_in_bytes']:
        p = Path(candidate)
        if p.exists():
            value = p.read_text().strip()
            if value.isdigit() and int(value) < 2 ** 60: env['container_RAM_limit_GiB'] = int(value) / 1024 ** 3
    summary['environment'] = env
    summary['benchmarks'] = {'csv_MB': csv_path.stat().st_size / 1e6, 'parquet_MB': parquet_path.stat().st_size / 1e6, 'csv_read_s': tcsv, 'parquet_read_s': tpq, 'duckdb_aggregation_s': td, 'pandas_parquet_aggregation_s': tp, 'csv_runs': csv_runs, 'parquet_runs': pq_runs, 'duckdb_runs': duck_runs, 'pandas_aggregation_runs': pandas_agg_runs, 'protocol': 'Медиана 3 последовательных запусков; кэш ОС не очищался. В тесте чтения читаются все столбцы. Для сравнения агрегаций pandas и DuckDB используют один Parquet и только arxiv_month, citationCount. DuckDB: 1 поток. Это среда выполнения, не компьютер студентки.'}
    save_table(pd.DataFrame([{'metric': k, 'value': v} for k, v in summary['benchmarks'].items() if not isinstance(v, list)]), root / 'results/benchmarks.csv')
    write_json(root / 'results/environment.json', env)
    chart_specs = []
    def chart(name, title, note, draw):
        fig, ax = plt.subplots(figsize=(9.6, 5.2))
        draw(ax)
        ax.set_title(title, fontsize=15, pad=12)
        ax.tick_params(labelsize=10)
        fig.tight_layout()
        fig.savefig(root / 'results/figures' / (name + '.png'), dpi=170)
        plt.close(fig)
        chart_specs.append({'file': name + '.png', 'title': title, 'note': note})
    q = quality.sort_values('missing_pct')
    chart('01_missing', 'Пропуски в исходных полях', f'Аннотация отсутствует у {show_num(summary["missing_abstract_raw"])} исходных записей. Отсутствующие тексты не придумываются: используется название. Нули на графике означают отсутствие обнаруженных пропусков.', lambda ax: (ax.barh(q.field, q.missing_pct), ax.set_xlabel('Доля пропусков, %'), ax.set_xlim(0, max(1, float(q.missing_pct.max()) * 1.12))))
    mc = clean.arxiv_month.value_counts().sort_index()
    chart('02_month', 'Покрытие по месяцу в arXiv ID', f'Месяц извлечён из arXiv ID, а не из publicationDate. Отсутствующих месяцев 2025 года: {len(summary["missing_arxiv_months_2025"])}; перечень сохранён в run_summary.json. Это покрытие поднабора, не статистика всего arXiv.', lambda ax: (ax.bar(mc.index, mc.values), ax.tick_params(axis='x', rotation=60), ax.set_ylabel('Публикации')))
    yc = clean.year.dropna().astype(int).value_counts().sort_index()
    chart('03_year', 'Фактические значения поля year', f'{show_num(summary["year_not_2025"])} записей имеют заполненный year, отличный от 2025. Название папки не подменяет дату публикации. Причина расхождения требует сверки с первоисточниками.', lambda ax: (ax.bar(yc.index.astype(str), yc.values), ax.set_ylabel('Публикации'), ax.set_xlabel('year из файла')))
    citations = clean.citationCount.dropna().astype(float)
    chart('04_citations', 'Распределение цитирований', f'Медиана: {show_num(summary["citation_median"], 1)}; максимум: {show_num(summary["citation_max"])}. Использовано log10(1 + число цитирований), чтобы показать и малые, и большие значения. Цитируемость не является меткой релевантности.', lambda ax: (ax.hist(np.log10(1 + citations), bins=45), ax.set_xlabel('log10(1 + citationCount)'), ax.set_ylabel('Публикации')))
    aw = clean.loc[clean.has_abstract, 'abstract_words']
    limit = max(1, int(aw.quantile(.99))) if len(aw) else 1
    chart('05_abstract', 'Длина доступных аннотаций', f'Медиана: {show_num(summary["abstract_words_median"])} слов. Показаны длины до 99-го процентиля ({limit} слов); более длинные тексты сохранены в данных. Пустые аннотации в эту гистограмму не входят.', lambda ax: (ax.hist(aw[aw <= limit], bins=40), ax.set_xlabel('Слова, разделённые пробелами'), ax.set_ylabel('Публикации')))
    ac = clean.author_count
    alimit = max(1, int(ac.quantile(.99)))
    chart('06_authors', 'Количество авторов в записи', f'Медиана: {show_num(summary["author_median"])}. Счётчик получен разделением authors по точке с запятой. На графике значения до 99-го процентиля ({alimit}); ноль означает отсутствие списка, а не доказанное отсутствие авторов.', lambda ax: (ax.hist(ac[ac <= alimit], bins=min(alimit + 1, 40)), ax.set_xlabel('Число имён в authors'), ax.set_ylabel('Публикации')))
    field_series = clean.fieldsOfStudy.fillna('').map(lambda x: sorted({v.strip() for v in x.split(';') if v.strip()})).explode().dropna()
    field_counts = field_series.value_counts()
    top = field_counts.drop('Computer Science', errors='ignore').head(10).sort_values()
    summary['field_count_unique'] = len(field_counts)
    summary['multiple_fields_rows'] = int((clean.field_count > 1).sum())
    def field_plot(ax):
        if len(top): ax.barh(top.index, top.values)
        else: ax.text(.5, .5, 'Дополнительных областей не обнаружено', ha='center', va='center', transform=ax.transAxes)
        ax.set_xlabel('Публикации с указанной областью')
    chart('07_fields', 'Дополнительные научные области', f'Найдено {len(field_counts)} отдельных обозначений областей. В {show_num(summary["multiple_fields_rows"])} записях несколько областей. Computer Science исключена только с этого графика; одна статья может учитываться в нескольких столбцах.', field_plot)
    pair = clean[['author_count', 'citationCount']].dropna().astype(float)
    corr = pair.corr(method='spearman').iloc[0, 1] if len(pair) > 1 else np.nan
    summary['author_citation_spearman'] = float(corr) if np.isfinite(corr) else None
    points = pair.sample(min(5000, len(pair)), random_state=42)
    chart('08_relation', 'Авторы и цитирования: связь признаков', f'Корреляция Спирмена на доступных парах: {show_num(summary["author_citation_spearman"], 3)}. Для читаемости показано до 5 000 случайных точек, seed=42. Связь не доказывает причинность; популярность не используется в поисковом score.', lambda ax: (ax.scatter(points.author_count, np.log10(1 + points.citationCount), s=8, alpha=.25), ax.set_xlabel('Количество авторов'), ax.set_ylabel('log10(1 + citationCount)')))
    summary['charts'] = chart_specs
    print('Графики построены. Строю разреженный TF-IDF индекс...', flush=True)
    texts = clean.title.fillna('') + ' ' + clean.abstract.fillna('')
    vectorizer = TfidfVectorizer(lowercase=True, stop_words='english', min_df=2, max_df=.98, max_features=80000, ngram_range=(1, 1), sublinear_tf=True, norm='l2', dtype=np.float32)
    started = time.perf_counter()
    matrix = vectorizer.fit_transform(texts)
    summary['index_seconds'] = time.perf_counter() - started
    summary['index_shape'] = list(matrix.shape)
    summary['index_nnz'] = int(matrix.nnz)
    summary['index_sparse_MB'] = (matrix.data.nbytes + matrix.indices.nbytes + matrix.indptr.nbytes) / 1e6
    summary['zero_vector_rows'] = int((matrix.getnnz(axis=1) == 0).sum())
    sparse.save_npz(root / 'index/tfidf.npz', matrix)
    joblib.dump(vectorizer, root / 'index/vectorizer.joblib', compress=3)
    clean[['arxiv_id', 'title', 'abstract', 'authors', 'year']].to_parquet(root / 'index/documents.parquet', index=False)
    def rank(query, k=10):
        qv = vectorizer.transform([query])
        if qv.nnz == 0: return [], np.zeros(len(clean), dtype=np.float32)
        scores = (matrix @ qv.T).toarray().ravel()
        order = np.argsort(-scores, kind='stable')
        order = order[scores[order] > 0][:k]
        return order.tolist(), scores
    answers, query_times = [], []
    for qid, query, ru in QUERIES:
        started = time.perf_counter()
        order, scores = rank(query)
        query_times.append(time.perf_counter() - started)
        for pos, idx in enumerate(order, 1):
            row = clean.iloc[idx]
            answers.append({'query_id': qid, 'query': query, 'query_ru': ru, 'rank': pos, 'arxiv_id': row.arxiv_id, 'title': row.title, 'abstract': '' if pd.isna(row.abstract) else row.abstract, 'score': float(scores[idx]), 'url': 'https://arxiv.org/abs/' + row.arxiv_id, 'source_sha256': digest.hexdigest()})
    answer_columns = ['query_id', 'query', 'query_ru', 'rank', 'arxiv_id', 'title', 'abstract', 'score', 'url', 'source_sha256']
    results = pd.DataFrame(answers, columns=answer_columns)
    save_table(results, root / 'results/search_results.csv')
    save_table(pd.DataFrame(QUERIES, columns=['query_id', 'query', 'query_ru']), root / 'evaluation/queries.csv')
    labels = results.copy()
    labels['relevant'], labels['reviewer'], labels['reason'] = '', '', ''
    labels_path = root / 'evaluation/relevance_labels.csv'
    if not labels_path.exists(): save_table(labels, labels_path)
    summary['query_time_median_s'] = median(query_times)
    summary['returned_results'] = len(results)
    summary['query_times_s'] = query_times
    # Это только техническая проверка поиска известной записи, не тест смысловой релевантности.
    eligible = np.flatnonzero(matrix.getnnz(axis=1) > 0)
    sample_ids = np.random.default_rng(42).choice(eligible, size=min(100, len(eligible)), replace=False)
    tests = []
    for idx in sample_ids:
        order, _ = rank(str(clean.iloc[idx].title))
        rank_pos = order.index(int(idx)) + 1 if int(idx) in order else None
        tests.append({'arxiv_id': clean.iloc[idx].arxiv_id, 'query_is_exact_title': clean.iloc[idx].title, 'rank_of_same_record': rank_pos})
    test_df = pd.DataFrame(tests)
    save_table(test_df, root / 'results/known_item_tests.csv')
    summary['known_item_n'] = len(tests)
    summary['known_item_hit1'] = sum(t['rank_of_same_record'] == 1 for t in tests) / len(tests) if tests else None
    summary['known_item_hit10'] = sum(t['rank_of_same_record'] is not None for t in tests) / len(tests) if tests else None
    summary['known_item_warning'] = 'Запрос равен заголовку записи, которая есть в индексе. Это проверка целостности реализации, не независимая оценка пользовательского поиска.'
    summary['evaluation'] = evaluate_labels(root, summary)
    html_parts = ['<!doctype html><html lang="ru"><meta charset="utf-8"><title>Проверка выдачи</title><style>body{font:17px system-ui;max-width:1000px;margin:40px auto;line-height:1.55}article{border:1px solid #ccc;padding:20px;margin:18px 0}h1,h2{line-height:1.2}small{display:block}textarea{width:98%;height:60px}</style><h1>Сохранённые результаты поиска</h1><p>Это не работающий веб-агент. Здесь сохранена выдача 10 запросов для проверки. Откройте CSV relevance_labels.csv и заполните relevant (1 или 0), reviewer и reason. Релевантность определяется соответствием конкретному запросу, а не количеством цитирований.</p>']
    for qid, query, ru in QUERIES:
        html_parts.append('<h2>' + html.escape(qid + ': ' + ru) + '</h2><p>' + html.escape(query) + '</p>')
        for _, row in results[results.query_id == qid].iterrows():
            html_parts.append('<article><b>' + str(row['rank']) + '. ' + html.escape(row.title) + '</b><small>' + html.escape(row.arxiv_id) + '; score=' + f'{row.score:.4f}' + '</small><p>' + html.escape(row.abstract) + '</p><a href="' + html.escape(row.url, quote=True) + '" target="_blank" rel="noopener">Открыть источник</a></article>')
    html_parts.append('</html>')
    (root / 'evaluation/review_results.html').write_text('\n'.join(html_parts), encoding='utf-8')
    make_review_form(root, results, digest.hexdigest())
    checklist = [
        ('Объём и структура', f'{n} исходных строк; {len(raw.columns)} полей; {len(clean)} строк после очистки.'),
        ('Гранулярность', 'Одна строка — одна запись публикации; рабочий ключ arxiv_id. Повторы нормализованных заголовков отмечены как кандидаты на дополнительную проверку, не объявлены доказанными дублями.'),
        ('Дубли', f'Лишних строк по нормализованному ключу до очистки: {summary["duplicate_id_extra_rows"]}; удалено после проверки: {summary["duplicates_removed"]}.'),
        ('Пропуски', 'Проценты для каждого поля сохранены в missing_values.csv. Нет аннотации — поиск по названию; неизвестные цитирования не заменены на ноль.'),
        ('Числовые значения', 'Минимум, медиана, максимум и число некорректных значений сохранены в numeric_quality.csv.'),
        ('Временное покрытие', f'publicationDate: {summary["min_publication_date"]} — {summary["max_publication_date"]}; year != 2025: {summary["year_not_2025"]}; несогласованность year и даты: {summary["year_date_mismatches"]}.'),
        ('Категории', f'{len(field_counts)} отдельных названий областей после разделения по ;. company сохранено, но не интерпретируется как достоверная принадлежность автора.'),
        ('Целевая переменная', 'Корпус не содержит меток релевантности для пользовательских запросов. Подготовлена ручная разметка. Precision@10 до её заполнения не объявляется.')
    ]
    save_table(pd.DataFrame(checklist, columns=['check', 'result']), root / 'results/quality_checklist.csv')
    summary['checklist'] = checklist
    source_file = Path(__file__).resolve()
    if source_file != (root / 'analysis.py').resolve(): shutil.copyfile(source_file, root / 'analysis.py')
    dependencies = ['numpy', 'pandas', 'scipy', 'scikit-learn', 'matplotlib', 'pyarrow', 'duckdb', 'joblib', 'psutil', 'python-pptx', 'python-docx']
    (root / 'requirements.txt').write_text('\n'.join(name + '==' + importlib.metadata.version(name) for name in dependencies) + '\n', encoding='utf-8')
    (root / '.gitignore').write_text('data/raw/*\ndata/processed/*\ndata/features/*\nindex/*\n__pycache__/\n.venv/\n*.zip\n', encoding='utf-8')
    write_json(root / 'results/run_summary.json', summary)
    (root / 'data/raw/README.md').write_text('Исходный файл metadata.jsonl сохраняется неизменённым. SHA-256: ' + digest.hexdigest() + '\nИсточник: ' + DATA_URL + '\nПоднабор: Computer_Science/2025. В лёгкий архив исходник не включён; используйте свой скачанный файл.\n', encoding='utf-8')
    return summary




def make_review_form(root, results, sha):
    data = results.fillna('').to_dict(orient='records')
    encoded = json.dumps(data, ensure_ascii=False).replace('<', '\\u003c')
    opening = '''<!doctype html><html lang="ru"><meta charset="utf-8"><title>Проверка релевантности</title><style>body{font:17px system-ui;line-height:1.55;max-width:1050px;margin:30px auto;padding:0 18px}header{position:sticky;top:0;background:white;border-bottom:2px solid #ccc;padding:14px;z-index:2}article{border:1px solid #bbb;margin:20px 0;padding:20px;border-radius:8px}label{margin-right:25px}textarea{width:97%;min-height:65px;margin-top:12px}button,input{font:inherit;padding:8px}small{display:block}h2{line-height:1.3}</style><header><b>Ручная проверка поиска</b><br><label>Проверяющий: <input id="reviewer" placeholder="Имя и фамилия"></label><button id="save">Сохранить CSV</button> <span id="progress"></span></header><h1>Соответствует ли статья запросу?</h1><p>Прочитайте название и аннотацию. Выберите 1, если статья содержательно отвечает намерению запроса, и 0 — если нет. Одного совпадения слов или высокой цитируемости недостаточно. Объяснение обязательно. Оценка относится только к видимым метаданным, не к полному тексту. Ответы автоматически сохраняются в этом браузере; никуда не отправляются.</p><div id="items"></div><script>const DATA='''
    js = r''';
const key='kamilla-relevance-'+(DATA[0]?.source_sha256 || 'empty');
let saved={}; try { saved=JSON.parse(localStorage.getItem(key)||'{}'); } catch(e) {}
const reviewer=document.getElementById('reviewer'); reviewer.value=saved.reviewer||'';
const host=document.getElementById('items'); let last='';
function remember(){const rows=DATA.map((d,i)=>({relevant:document.querySelector('input[name="r'+i+'"]:checked')?.value||'',reason:document.getElementById('why'+i).value}));try{localStorage.setItem(key,JSON.stringify({reviewer:reviewer.value,rows:rows}));}catch(e){}document.getElementById('progress').textContent=rows.filter(r=>r.relevant!==''&&r.reason.trim()).length+' / '+DATA.length;return rows;}
DATA.forEach((d,i)=>{if(d.query_id!==last){let h=document.createElement('h2');h.textContent=d.query_id+': '+d.query_ru+' / '+d.query;host.appendChild(h);last=d.query_id;}let a=document.createElement('article');let h=document.createElement('h3');h.textContent=d.rank+'. '+d.title;a.appendChild(h);let meta=document.createElement('small');meta.textContent='arXiv '+d.arxiv_id+'; сходство '+Number(d.score).toFixed(4)+' (не вероятность)';a.appendChild(meta);let p=document.createElement('p');p.textContent=d.abstract||'Аннотация отсутствует: оценка только по названию, при сомнении проверьте источник.';a.appendChild(p);let link=document.createElement('a');link.textContent='Открыть статью';link.href=d.url;link.target='_blank';link.rel='noopener';a.appendChild(link);a.appendChild(document.createElement('br'));['1','0'].forEach(value=>{let label=document.createElement('label');let radio=document.createElement('input');radio.type='radio';radio.name='r'+i;radio.value=value;radio.checked=saved.rows?.[i]?.relevant===value;radio.addEventListener('change',remember);label.appendChild(radio);label.appendChild(document.createTextNode(value==='1'?'Подходит (1)':'Не подходит (0)'));a.appendChild(label);});let reason=document.createElement('textarea');reason.id='why'+i;reason.placeholder='Почему статья подходит или не подходит этому запросу?';reason.value=saved.rows?.[i]?.reason||'';reason.addEventListener('input',remember);a.appendChild(reason);host.appendChild(a);});
reviewer.addEventListener('input',remember);remember();
document.getElementById('save').onclick=()=>{const rows=remember();if(!reviewer.value.trim()){alert('Введите имя проверяющего.');return;}if(rows.some(r=>r.relevant===''||!r.reason.trim())&&!confirm('Часть оценок не заполнена. Сохранить промежуточный файл? Метрики пока не будут рассчитаны.'))return;const headers=['query_id','query','query_ru','rank','arxiv_id','title','abstract','score','url','source_sha256','relevant','reviewer','reason'];const safe=v=>'"'+String(v??'').replace(/"/g,'""')+'"';let csv=headers.map(safe).join(',')+'\r\n';DATA.forEach((d,i)=>{let row={...d,...rows[i],reviewer:reviewer.value.trim()};csv+=headers.map(h=>safe(row[h])).join(',')+'\r\n';});let a=document.createElement('a');const url=URL.createObjectURL(new Blob(['\ufeff'+csv],{type:'text/csv;charset=utf-8'}));a.href=url;a.download='relevance_labels.csv';a.click();setTimeout(()=>URL.revokeObjectURL(url),3000);};
</script></html>'''
    (root / 'evaluation/review_form.html').write_text(opening + encoded + js, encoding='utf-8')




def make_materials(root, summary):
    s = summary
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    notes = []
    ink, muted, accent, pale = RGBColor(24, 41, 63), RGBColor(83, 98, 116), RGBColor(39, 99, 161), RGBColor(242, 246, 250)
    def text(slide, txt, x, y, w, h, size=22, bold=False, color=None):
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(.04)
        for i, line in enumerate(str(txt).split('\n')):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = line
            p.font.name = 'Arial'
            p.font.size = Pt(size)
            p.font.bold = bold
            p.font.color.rgb = color or ink
            p.space_after = Pt(10)
        return box
    def new_slide(title, speech, source='Расчёты: results/run_summary.json; исходник metadata.jsonl'):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(.12))
        bar.fill.solid(); bar.fill.fore_color.rgb = accent; bar.line.fill.background()
        text(slide, title, .55, .35, 12.2, 1.65 if '\n' in title else .82, 29, True)
        text(slide, f'{AUTHOR}  |  Аналитика больших данных в бизнесе  |  {len(prs.slides):02}', .55, 7.02, 12.1, .2, 10, color=muted)
        text(slide, source, .55, 6.7, 12.1, .24, 10, color=muted)
        slide.notes_slide.notes_text_frame.text = speech + '\n\nИсточник: ' + source
        notes.append((title, speech))
        return slide
    def body(title, paragraphs, speech, source='Проектная постановка; исправленная заявка, занятие 2'):
        slide = new_slide(title, speech, source)
        text(slide, '\n'.join(paragraphs), .75, 1.55, 11.7, 4.9, 23)
        return slide
    def box(slide, label, x, y, w=2.1, h=.75, kind=MSO_SHAPE.ROUNDED_RECTANGLE):
        shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
        shape.fill.solid(); shape.fill.fore_color.rgb = pale
        shape.line.color.rgb = accent
        tf = shape.text_frame; tf.word_wrap = True
        tf.margin_top = Inches(.08); tf.margin_bottom = Inches(.02)
        tf.text = label
        for p in tf.paragraphs:
            p.font.name = 'Arial'; p.font.size = Pt(15); p.font.color.rgb = ink
        return shape
    def arrow(slide, x1, y1, x2, y2):
        line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        line.line.color.rgb = muted; line.line.width = Pt(1.6)
        tail = OxmlElement('a:tailEnd'); tail.set('type', 'triangle')
        line._element.spPr.get_or_add_ln().append(tail)
    slide = new_slide('Разработка ИИ-агента\nдля преподавателя, исследователя', 'Здравствуйте. Меня зовут Морбекова Камилла. Я представляю аналитическую часть дипломной темы: поиск научных публикаций. В этой работе я не заявляю разработку всего агента. Я исследую пригодность данных и базовый способ ранжирования источников.', 'Тема диплома № 13; автор: Морбекова Камилла')
    text(slide, 'Аналитическая часть: поиск научных публикаций', .75, 2.6, 11.6, .8, 29, True)
    text(slide, f'{AUTHOR}\nПервый рубежный контроль\nКорпус после очистки: {show_num(s["clean_rows"])} публикаций', .75, 4, 11, 1.7, 24)
    body('Что решается и кто принимает решение', ['Пользователь: преподаватель или исследователь.', 'Ввод: тема на английском языке. Результат: Top-10 источников, а не готовый обзор.', 'Человек выбирает, какие статьи читать и использовать.', 'Лишняя статья тратит время; пропущенная важная статья обедняет обзор.', 'Цель: не менее 7 подходящих статей в Top-10; это проверяемая цель, не заранее полученный результат.'], 'Пользователь ищет литературу для лекции или исследования. Модель только упорядочивает найденные источники. Решение остаётся у человека. Есть два вида ошибок: показать неподходящую статью и не показать нужную. Числовую стоимость ошибок мы не измеряли и не придумываем.')
    body('Какие данные загружены', [f'Исходник: {show_num(s["raw_rows"])} записей, {s["raw_columns"]} полей, {show_num(s["raw_jsonl_MB"], 2)} МБ JSONL.', 'Одна запись — публикация. Ключ: arxiv_id.', 'Файл: Computer_Science/2025; фактический year проверяется отдельно.', 'Название + аннотация — вход поиска. Цитирования — только для анализа.', 'Происхождение: arXiv + Semantic Scholar; полученный файл — один обогащённый источник.'], 'JSONL означает, что каждая непустая строка содержит отдельный JSON-объект. Мы используем предоставленный снимок данных и сохраняем его контрольную сумму. Важно: имя папки 2025 не гарантирует, что поле year в каждой записи равно 2025.', DATA_URL + '; results/run_summary.json')
    body('Восемь проверок качества: результат', [f'Структура: {show_num(s["raw_rows"])} строк → {show_num(s["clean_rows"])} после очистки.', f'Ключ: {s["invalid_or_missing_id"]} некорректных/пустых ID; {s["duplicates_removed"]} удалённых повторов.', f'Тексты: {show_num(s["missing_title_raw"])} пустых названий; {show_num(s["abstract_missing_clean"])} аннотаций без текста после очистки.', f'Числа: {s["invalid_numeric_count"]} некорректных значений; даты: {s["invalid_dates"]} ошибок разбора.', f'Период: {show_num(s["year_not_2025"])} заполненных year ≠ 2025; областей: {s["field_count_unique"]}.', 'Готовых меток релевантности нет: предусмотрена отдельная ручная проверка.'], 'Все восемь пунктов подробно сохранены в quality_checklist.csv. Отдельные таблицы содержат пропуски по полям и минимумы, медианы, максимумы чисел. Смысл проверки не в поиске красивых цифр, а в понимании ограничений данных.', 'results/quality_checklist.csv; missing_values.csv; numeric_quality.csv')
    body('Подготовка: исходник не меняем', ['raw — исходный JSONL и его SHA-256.', 'processed — ключи, типы, обработка повторов; CSV и Parquet одного содержания.', 'features — число авторов, областей, слов; агрегаты по месяцу arXiv ID.', f'Проверка соединения: {show_num(s["join_rows_before"])} → {show_num(s["join_rows_after"])} строк.', 'Нет аннотации — используем название. Неизвестные цитирования не равны нулю.', 'year и publicationDate не переписываются под название папки.'], 'Подготовка разделена на три слоя, чтобы обработку можно было повторить. Для повторяющегося ID оставляем запись с более длинной аннотацией, при равенстве — более раннюю исходную строку. Это явно заданное правило. Соединение с месячными агрегатами проверено как many-to-one. Оно не является вторым независимым источником.', 'analysis.py: run_analysis; data/raw, processed, features')
    slide = new_slide('BPMN: от запроса до выбора источников', 'На схеме один процесс и две дорожки: человек и поисковый модуль. Человек задаёт тему. Модуль рассчитывает выдачу. Человек оценивает источники. Если они не подходят, запрос уточняется. Вывод схемы: финальное решение нельзя передавать алгоритму.', 'Развитие схемы из занятия 3; проектная BPMN-схема')
    for y, label in [(1.5, 'Преподаватель / исследователь'), (3.5, 'Поисковый модуль')]:
        lane = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.55), Inches(y), Inches(12.15), Inches(1.8)); lane.fill.background(); lane.line.color.rgb = muted
        text(slide, label, .7, y + .05, 5, .3, 13, True)
    box(slide, '', .8, 2.05, .7, .7, MSO_SHAPE.OVAL)
    text(slide, 'Нужны источники', .65, 2.85, 1.2, .55, 11)
    box(slide, 'Ввести / уточнить\nзапрос', 2, 2, 2, .8)
    box(slide, 'Векторизация\nзапроса', 3.6, 4.05, 2, .8)
    box(slide, 'Рассчитать\nсходство', 6, 4.05, 2, .8)
    box(slide, 'Выдать Top-10', 8.5, 4.05, 2, .8)
    box(slide, 'Просмотреть\nисточники', 8.3, 2, 2, .8)
    box(slide, 'X', 10.65, 1.95, 1, .95, MSO_SHAPE.DIAMOND)
    text(slide, 'Подходит?', 10.5, 1.55, 1.3, .35, 12)
    end_event = box(slide, '', 11.9, 2.08, .7, .7, MSO_SHAPE.OVAL)
    end_event.line.width = Pt(3)
    text(slide, 'Выбор сделан', 11.7, 2.85, .9, .55, 11)
    for coords in [(1.55, 2.4, 2, 2.4), (4, 2.5, 4.6, 4.05), (5.6, 4.45, 6, 4.45), (8, 4.45, 8.5, 4.45), (9.5, 4.05, 9.3, 2.8), (10.3, 2.4, 10.65, 2.4), (11.65, 2.4, 11.9, 2.4), (11.15, 2.9, 11.15, 3.25), (11.15, 3.25, 3, 3.25), (3, 3.25, 3, 2.8)]: arrow(slide, *coords)
    text(slide, 'Да', 11.65, 1.7, .6, .3, 12)
    text(slide, 'Нет: уточнить запрос', 5.2, 2.95, 2.6, .35, 13)
    text(slide, 'Вывод: алгоритм ранжирует; окончательный выбор делает человек.', .75, 5.7, 11.8, .7, 22, True)
    slide = new_slide('Пайплайн: подготовка отдельно от поиска', 'Сначала исходный файл очищается и превращается в поисковый индекс. Это делается один раз для выбранного снимка данных. Затем каждый запрос преобразуется тем же словарём. Вычисляется сходство с документами, а не матрица всех пар публикаций. Вывод: подготовку корпуса не нужно повторять при каждом запросе.', 'analysis.py; index/; схема пайплайна — вторая из двух схем')
    for label, x in [('JSONL\nraw', .8), ('Проверки\nи очистка', 3.25), ('title + abstract\nfeatures', 5.7), ('TF-IDF\nиндекс', 8.15)]: box(slide, label, x, 1.85, 2.05, 1)
    for x in [2.85, 5.3, 7.75]: arrow(slide, x, 2.35, x + .4, 2.35)
    for label, x in [('Запрос\nпользователя', .8), ('Тот же\nсловарь TF-IDF', 3.25), ('Косинусное\nсходство', 5.7), ('Top-10\nпубликаций', 8.15)]: box(slide, label, x, 4.1, 2.05, 1)
    for x in [2.85, 5.3, 7.75]: arrow(slide, x, 4.6, x + .4, 4.6)
    arrow(slide, 9.2, 2.85, 6.7, 4.1)
    text(slide, 'Вывод: индекс строится заранее; для нового запроса пересчитывается только выдача.', .75, 5.8, 11.8, .7, 22, True)
    b = s['benchmarks']; e = s['environment']
    bench_text = [f'CSV: {show_num(b["csv_MB"], 2)} МБ  |  Parquet: {show_num(b["parquet_MB"], 2)} МБ', f'Чтение CSV: {show_num(b["csv_read_s"], 4)} с  |  Parquet: {show_num(b["parquet_read_s"], 4)} с', f'Агрегация по одному Parquet: DuckDB {show_num(b["duckdb_aggregation_s"], 4)} с; pandas {show_num(b["pandas_parquet_aggregation_s"], 4)} с', f'ОЗУ среды: {show_num(e["system_RAM_GiB"], 2)} ГиБ; Python {e["python"]}.', 'Медиана трёх запусков; кэш ОС не очищался; значения не универсальны.', 'Выбор для этого запуска: pandas для подготовки, Parquet для хранения, DuckDB для SQL-агрегаций.']
    body('Замеры хранения и обработки', bench_text, 'Здесь реальные результаты данного запуска, а не числа из чужого примера. CSV и Parquet содержат одинаковую таблицу, совпадение проверено. Агрегации проверены на одинаковый результат. Для обработки этого снимка достаточно использованной однокомпьютерной среды; это не доказывает пригодность для любых объёмов.', 'results/benchmarks.csv; results/environment.json; protocol в run_summary.json')
    for spec in s['charts']:
        slide = new_slide(spec['title'], spec['note'], 'analysis.py; results/figures/' + spec['file'])
        slide.shapes.add_picture(str(root / 'results/figures' / spec['file']), Inches(.5), Inches(1.35), width=Inches(9.05))
        text(slide, 'ВЫВОД', 9.85, 1.65, 2.7, .4, 17, True, accent)
        text(slide, spec['note'], 9.8, 2.25, 2.9, 4.05, 17)
    body('Baseline: TF-IDF + косинусное сходство', ['Текст документа = название + доступная аннотация.', 'TF-IDF выделяет значимые слова; это лексический, не полноценный семантический поиск.', f'Индекс: {show_num(s["index_shape"][0])} × {show_num(s["index_shape"][1])}; разреженное хранение {show_num(s["index_sparse_MB"], 1)} МБ.', 'Параметры: min_df=2, max_df=0,98, до 80 000 признаков; seed=42 в выборках.', f'Построение индекса: {show_num(s["index_seconds"], 2)} с; медиана запроса: {show_num(s["query_time_median_s"], 4)} с.', 'Score — сходство, а не вероятность правильности. Цитирования в score не входят.'], 'TF-IDF даёт больший вес словам, которые полезны для различения документов. После нормировки скалярное произведение даёт косинусное сходство. Обучаемого классификатора и меток 0/1 для обучения здесь нет. Индекс строится на доступном корпусе; качество пользовательских запросов нужно проверять отдельно.', SOURCES[2][1] + '; SOURCES: cosine_similarity; analysis.py')
    results = pd.read_csv(root / 'results/search_results.csv', dtype={'arxiv_id': 'string'}, keep_default_na=False)
    example = results[results.query_id == 'Q01'].head(3)
    ex_paragraphs = ['Запрос: large language models in education']
    for _, row in example.iterrows():
        title = row.title if len(row.title) <= 150 else row.title[:147] + '…'
        ex_paragraphs.append(f'{int(row["rank"])}. {title}\n    arXiv {row.arxiv_id}; score {row.score:.4f}')
    if not len(example): ex_paragraphs.append('Для запроса не найдено положительного совпадения словаря.')
    ex_paragraphs.append('Полные Top-10 и аннотации: search_results.csv / review_results.html.')
    example_slide = body('Пример фактической выдачи', ex_paragraphs, 'Это сохранённая выдача алгоритма для первого заранее заданного запроса. Порядок не исправлялся вручную. Высокий score ещё не доказывает, что статья полезна преподавателю. Нужно прочитать название и аннотацию и сравнить с намерением запроса.', 'results/search_results.csv; query_id=Q01')
    for shape in example_slide.shapes:
        if shape.has_text_frame and shape.top == Inches(1.55):
            for p in shape.text_frame.paragraphs: p.font.size = Pt(19)
    body('Техническая проверка реализации', [f'Проверено {s["known_item_n"]} случайных записей; seed=42.', 'Запросом служило точное название уже имеющейся публикации.', f'Та же запись на первом месте: {show_num((s["known_item_hit1"] or 0) * 100, 1)}%.', f'Та же запись в Top-10: {show_num((s["known_item_hit10"] or 0) * 100, 1)}%.', 'Это проверка индекса и ранжирования — НЕ Precision@10 пользовательских запросов.', 'Тест не доказывает понимание смысла, полноту поиска или пользу для преподавателя.'], 'Мы разделяем техническую исправность и полезность. Поиск по точному заголовку — упрощённый тест, ведь этот заголовок уже находится в индексе. Полученные проценты нельзя выдавать за точность рекомендаций. Полезность требует другой процедуры: независимой ручной оценки тематических запросов.', 'results/known_item_tests.csv; предупреждение known_item_warning')
    ev = s['evaluation']
    if ev['status'] == 'complete':
        eval_par = [f'Ручная проверка: {ev["queries"]} запросов.', f'Средняя Precision@10: {show_num(ev["mean_precision_at_10"], 3)}.', f'Цель ≥7 из 10 выполнена на {ev["queries_meeting_7_of_10"]} из {ev["queries"]} запросов.', 'Проверяющие: ' + ', '.join(ev['reviewers']), 'Это оценка по введённым меткам; она не измеряет Recall по всему корпусу.']
        eval_speech = 'Precision@10 — доля подходящих статей среди первых десяти. Метки внесены человеком и сохранены с пояснениями. Мы показываем среднее и число запросов, достигших цели. Уровень квалификации проверяющего не подтверждался отдельно; это нужно учитывать при интерпретации.'
    else:
        eval_par = ['Подготовлены 10 тематических запросов и их фактическая выдача.', 'Ручная оценка НЕ завершена. Precision@10 пока не рассчитана.', 'Для каждой пары: relevant=1/0, имя проверяющего и объяснение.', 'Precision@10 = число подходящих статей / 10.', 'Цель ≥7 из 10 пока НЕ подтверждена.', 'Нельзя подменять ручную проверку score модели или числом цитирований.']
        eval_speech = 'Самая важная незавершённая часть — оценка смыслового качества. Код и выдача есть, но человека нельзя заменить выдуманными экспертными метками. Поэтому цель семь из десяти пока не объявлена достигнутой. Файл для проверки подготовлен. После заполнения метрик презентация пересобирается без повторного анализа.'
    body('Оценка полезности: статус и критерий', eval_par, eval_speech, 'evaluation/relevance_labels.csv; results/relevance_metrics.csv при наличии')
    body('Ограничения и дальнейший план', ['Корпус ограничен выбранным поднабором; это не вся научная литература.', 'Различаются год папки, year и publicationDate: нужна сверка первоисточников.', 'Лексический поиск плохо учитывает синонимы и язык запроса.', 'Закончить ручную проверку, разобрать ошибки и проверить критерий 7/10.', 'Далее сравнить с BM25 / семантическим поиском на тех же запросах.', 'В дипломе: интеграция модуля в агента. Сам агент в этой работе не реализован.'], 'Ограничения являются частью результата. Мы не утверждаем, что весь диплом готов. Следующий шаг — завершить оценку полезности и только после неё сравнивать более сложные методы. Для сравнения нужен общий фиксированный набор запросов и одинаковые правила разметки.', 'Результаты анализа; план команды, не выполненные эксперименты')
    body('Воспроизводимость и оставшиеся требования', ['Код: analysis.py; версии: requirements.txt; снимок данных: SHA-256.', 'Чистый запуск пересобирает данные, графики, индекс и материалы.', 'Повторные замеры скорости могут отличаться; содержание снимка проверяется хешем.', 'Внешних исходных систем две, но загруженный обогащённый файл один.', 'Самостоятельное соединение двух независимых источников НЕ выполнено; нужно согласование преподавателя.', 'Перед защитой: завершить разметку, открыть презентацию и проверить читаемость всех слайдов.'], 'Воспроизводимость означает, что известны исходные данные, версии и порядок запуска. В лёгком архиве нет большого корпуса — его нужно положить в data/raw. Отдельно остаётся вопрос требования двух источников: готовое обогащение мы не выдаём за выполненный нами внешний JOIN.', 'README.md; requirements.txt; results/run_summary.json; требования занятия 2')
    body('Источники', [name + '\n' + link for name, link in SOURCES[:5]], 'В работе используются предоставленный снимок корпуса и материалы курса. Методические сведения о TF-IDF, косинусном сходстве и формате идентификатора взяты из официальной документации. Полный список источников и оговорки о происхождении данных сохранены в README.', 'Атрибуция: yufan; arXiv; Semantic Scholar. Лицензия набора указана в карточке источника.')
    # Последний слайд содержит длинные ссылки — уменьшаем шрифт только этого блока.
    last = prs.slides[-1]
    for shape in last.shapes:
        if shape.has_text_frame and 'https://' in shape.text and shape.top > Inches(1):
            for p in shape.text_frame.paragraphs: p.font.size = Pt(16)
    ppt_path = root / 'materials/Kamilla_presentation.pptx'
    prs.save(ppt_path)
    # Проверка структуры файла, а не визуальный рендер.
    reloaded = Presentation(ppt_path)
    assert len(reloaded.slides) == len(notes)
    with zipfile.ZipFile(ppt_path) as archive: assert archive.testzip() is None
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Cm(1.8)
    normal = doc.styles['Normal']; normal.font.name = 'Arial'; normal.font.size = DocPt(11)
    doc.add_heading('Подготовка к защите с нуля', 0)
    doc.add_paragraph(AUTHOR + '\n' + DIPLOMA)
    doc.add_paragraph('Это пособие для выступления, а не отчёт по предмету. Презентация и расчёты сформированы запуском кода. Результаты ручной проверки нельзя объявлять полученными до заполнения разметки.')
    doc.add_heading('С чего начать', 1)
    doc.add_paragraph('Сначала прочитай постановку задачи. Затем открой презентацию и сопоставь каждый слайд с пояснением ниже. Для ручной проверки открой evaluation/review_form.html, укажи имя и отметь, подходит ли каждая статья запросу. Объясни оценку и нажми Сохранить CSV. Если непонятен английский текст, используй перевод, но проверяй именно соответствие запросу. Замени evaluation/relevance_labels.csv полученным файлом и пересобери материалы командой из README. Самостоятельный запуск полного анализа в чате не состоялся: эти материалы генерируются только в среде, где запущен код.')
    doc.add_heading('Главная мысль защиты', 1)
    doc.add_paragraph('Я исследую аналитическую часть функции поиска литературы для будущего ИИ-агента. Программа ранжирует публикации, а окончательное решение принимает преподаватель или исследователь. Я отдельно проверяю качество данных, техническую работу поиска и его полезность.')
    for i, (title, speech) in enumerate(notes, 1):
        doc.add_heading(f'Слайд {i}. {title.replace(chr(10), " ")}', 1)
        doc.add_paragraph(speech)
    glossary = [
        ('Датасет', 'Набор данных. Здесь это таблица записей о публикациях.'),
        ('Метаданные', 'Описание статьи: название, авторы, дата, аннотация. Это не обязательно полный текст статьи.'),
        ('JSONL', 'Текстовый формат: отдельный JSON-объект на каждой строке.'),
        ('CSV', 'Таблица в текстовом файле. При чтении программе приходится разбирать текст и определять типы.'),
        ('Parquet', 'Колоночный формат хранения. Типы сохраняются в файле; преимущество по размеру и времени проверяется измерениями, а не обещается заранее.'),
        ('pandas', 'Библиотека Python для работы с таблицами.'),
        ('DuckDB', 'Движок SQL-запросов к аналитическим данным, в том числе Parquet.'),
        ('Пропуск', 'Значение неизвестно или отсутствует. Это не то же самое, что ноль.'),
        ('Гранулярность', 'Что означает одна строка. Здесь одна строка соответствует записи публикации.'),
        ('Агрегация', 'Сводный расчёт: например, число статей и среднее число цитирований по месяцу.'),
        ('Baseline', 'Простой исходный способ решения, с которым затем сравнивают улучшения.'),
        ('TF-IDF', 'Числовое представление текста через веса слов. Важны частота слова в документе и его распространённость по корпусу.'),
        ('Разреженная матрица', 'Хранятся главным образом ненулевые веса. Не нужно создавать огромную плотную таблицу всех слов во всех документах.'),
        ('Косинусное сходство', 'Сравнение направлений числовых векторов текстов. Для ненулевых нормированных TF-IDF векторов вычисляется через скалярное произведение.'),
        ('Релевантность', 'Соответствие статьи конкретному запросу. Популярная статья не обязательно релевантна.'),
        ('Precision@10', 'Сколько из первых десяти результатов подходят запросу, делённое на десять.'),
        ('SHA-256', 'Контрольный отпечаток файла. Позволяет проверить, что использован тот же снимок данных.'),
        ('seed / random_state', 'Фиксированное начальное значение для повторяемой случайной выборки.')
    ]
    doc.add_heading('Словарь терминов', 1)
    for term, meaning in glossary:
        p = doc.add_paragraph(); p.add_run(term + '. ').bold = True; p.add_run(meaning)
    faq = [
        ('Где здесь ИИ-агент?', 'В дипломной теме. Здесь выполнен аналитический модуль одной его функции; агент целиком не реализован.'),
        ('Почему не 25 записей OpenAlex?', '25 записей относятся к прежнему упражнению по API. Для основного проекта в материалах курса требуется минимум 100 000 записей.'),
        ('Почему папка 2025, а встречается 2024?', 'Это обнаруженное расхождение описания и полей файла. Мы не знаем достоверную причину для каждой записи, не исправляем годы выдуманными значениями и отделяем месяц ID от publicationDate.'),
        ('Можно ли заменить пустые цитирования нулём?', 'Нет без обоснования: неизвестное число и реальное отсутствие цитирований — разные ситуации. В этом проекте пропуск сохраняется.'),
        ('Поиск умеет понимать смысл?', 'Baseline основан на словах и их весах. Он не гарантирует понимание синонимов и не является большим языковым агентом.'),
        ('Зачем проверка по точным названиям?', 'Это техническая проверка индекса. Она не является независимой оценкой релевантности и не заменяет ручную проверку запросов.'),
        ('Где train и test?', 'Это индексирование фиксированного корпуса, а не обучение классификатора на метках. Для оценки нужны отдельные запросы и разметка. Настройки нельзя подгонять под результаты окончательной проверки.'),
        ('Почему score не процент уверенности?', 'Score вычислен как сходство текстовых векторов. Он не откалиброван как вероятность полезности статьи.'),
        ('Почему не Spark?', 'Этот конкретный запуск обработал снимок на одной машине. Выбор объясняется измеренными объёмом, памятью и временем, а не модой на инструмент.'),
        ('Доказана ли цель 7 из 10?', 'Только после полной человеческой разметки. До этого в презентации явно написано, что цель не подтверждена.'),
        ('Проверены ли два источника?', 'У готового набора два заявленных происхождения: arXiv и Semantic Scholar. Самостоятельного внешнего JOIN мы не делали. Нужно уточнить, принимается ли готовое обогащение.'),
        ('Что именно делала студентка, а что ИИ?', 'Код и оформление подготовлены с помощью ИИ. На защите нужно честно назвать свою роль: запуск, проверка результатов, ручная разметка и объяснение решений — только если эти действия действительно выполнены.'),
        ('Все требования уже закрыты?', 'Нет, пока не завершена ручная оценка и не согласован вопрос двух источников. Презентация также требует просмотра после генерации; автоматическая проверка файла не заменяет визуальную проверку.')
    ]
    doc.add_heading('Вопросы преподавателя и ответы', 1)
    for question, answer in faq:
        doc.add_heading(question, 2); doc.add_paragraph(answer)
    doc.add_heading('Источники', 1)
    for name, link in SOURCES: doc.add_paragraph(name + ': ' + link)
    doc.save(root / 'materials/Kamilla_defense_guide.docx')
    (root / 'materials/speaker_notes.md').write_text('\n\n'.join(f'## {i}. {title}\n{speech}' for i, (title, speech) in enumerate(notes, 1)), encoding='utf-8')
    (root / 'README.md').write_text(f'''# Аналитический модуль поиска научных публикаций
Автор материалов: {AUTHOR}
Тема диплома № 13: {DIPLOMA}


## Статус
Расчёты этого комплекта сформированы запуском analysis.py. Ручная проверка: {s['evaluation']['status']}.
Критерий 7 релевантных из 10 не подтверждается до внесения человеческих меток.
Готовый файл объединяет заявленные данные arXiv и Semantic Scholar; самостоятельного внешнего JOIN нет.


## Запуск с нуля
1. Установите Python и зависимости: `python -m pip install -r requirements.txt`.
2. Положите свой metadata.jsonl в data/raw/metadata.jsonl.
3. Выполните `python analysis.py --input data/raw/metadata.jsonl --out .`.
Исходный файл не изменяется. Его SHA-256: {s['source_sha256']}.
Данные: {DATA_URL}, файл Computer_Science/2025/metadata.jsonl.
Новая загрузка с сайта может отличаться. Для точного повторения используйте исходный снимок с указанным SHA-256.


## Проверка полезности
Откройте evaluation/review_form.html в браузере. Введите имя проверяющего, отметьте Подходит/Не подходит и объясните оценку. Нажмите Сохранить CSV и замените evaluation/relevance_labels.csv полученным файлом. Сохранённая выдача без формы также есть в review_results.html.
Не ставьте 1 только из-за совпадения слов: название и аннотация должны отвечать намерению запроса.
Код не генерирует экспертные метки и не заменяет неизвестные метки на 0.
Затем: `python analysis.py --out . --refresh-presentation`.
Обновятся метрики, презентация и пособие. Формула: релевантные в Top-10 / 10; среднее — по 10 запросам.
Recall по всему корпусу не измеряется. Если результатов меньше десяти, незаполненные позиции не считаются релевантными.


## Демонстрация нового запроса
`python analysis.py --out . --query "scientific paper recommendation"`
Индекс должен быть предварительно построен. Не загружайте joblib-файлы из недоверенных источников.


## Структура
analysis.py — загрузка, аудит, подготовка, 8 графиков, замеры, TF-IDF, генерация PPTX и DOCX.
data/raw — неизменённый снимок; processed — очищенная таблица; features — производные признаки и агрегаты.
results — таблицы и графики; evaluation — запросы, выдача и ручные метки; index — поисковый индекс.
materials — презентация, речь и пособие.
requirements.txt фиксирует версии реально использованных библиотек. Для публикации в GitHub используйте лёгкий архив и .gitignore, а не большой корпус.


## Методологические ограничения
Год папки, год в ID и поле year не смешиваются. Метки company не используются как доказанные принадлежности авторов.
Числовые выбросы описываются, а не автоматически удаляются. Отсутствующие аннотации заменяются только отсутствием текста; поиск идёт по названию.
TF-IDF использует unigram, min_df=2, max_df=0.98, max_features=80000, L2, float32. Запросы не включены в обучение IDF.
Тест точного заголовка — технический sanity-check, не доказательство качества рекомендаций.
Замеры: медиана 3 запусков, кэш не сбрасывался; спецификация среды в results/environment.json.
Эти замеры относятся к среде выполнения, а не автоматически к личному компьютеру студентки.
После генерации откройте PPTX и DOCX и проверьте визуально. Код проверяет структуру PPTX, но не выполняет рендер всех страниц.


## Источники
''' + '\n'.join('- ' + name + ': ' + link for name, link in SOURCES), encoding='utf-8')
    status = f'''ГОТОВО ПО РЕЗУЛЬТАТАМ ЗАПУСКА: данные, 8 проверок, подготовка, замеры, 8 графиков, TF-IDF, технические тесты, презентация и пособие.
РУЧНАЯ РЕЛЕВАНТНОСТЬ: {s['evaluation']['status']}.
ТРЕБОВАНИЕ >=100000 ПОСЛЕ ОЧИСТКИ: {s['minimum_100k_satisfied']}.
ТРЕБУЕТ СОГЛАСОВАНИЯ: достаточно ли готового обогащения вместо самостоятельного соединения двух источников.
ПЕРЕД ЗАЩИТОЙ: заполнить ручную разметку; пересобрать материалы; открыть презентацию и проверить вёрстку; прочитать пособие и отрепетировать.
НЕ ЗАЯВЛЯТЬ: готовность всего агента, экспертную оценку без человека, точность рекомендаций по техническому тесту заголовков.
'''
    (root / 'START_HERE.txt').write_text(status, encoding='utf-8')
    archive_path = root.parent / (root.name + '_submission.zip')
    with zipfile.ZipFile(archive_path, 'w', compression=zipfile.ZIP_DEFLATED) as archive:
        for p in root.rglob('*'):
            if not p.is_file(): continue
            relative = p.relative_to(root)
            if relative.parts[0] == 'index': continue
            if relative.parts[0] == 'data' and p.suffix != '.md': continue
            if '__pycache__' in relative.parts: continue
            archive.write(p, arcname=str(Path(root.name) / relative))
    print('\nМатериалы сформированы:', flush=True)
    print(ppt_path, flush=True)
    print(root / 'materials/Kamilla_defense_guide.docx', flush=True)
    print(archive_path, flush=True)
    print('ВАЖНО: ' + s['evaluation']['message'], flush=True)
    return archive_path




def main():
    parser = argparse.ArgumentParser(description='Воспроизводимый анализ корпуса научных публикаций')
    parser.add_argument('--input', type=Path)
    parser.add_argument('--out', type=Path, default=Path('kamilla_project'))
    parser.add_argument('--refresh-presentation', action='store_true')
    parser.add_argument('--query', type=str)
    args = parser.parse_args()
    root = args.out.resolve()
    if args.query:
        vectorizer = joblib.load(root / 'index/vectorizer.joblib')
        matrix = sparse.load_npz(root / 'index/tfidf.npz')
        docs = pd.read_parquet(root / 'index/documents.parquet')
        qv = vectorizer.transform([args.query])
        if qv.nnz == 0:
            print('Нет слов запроса в словаре. Уточните английский запрос.'); return
        scores = (matrix @ qv.T).toarray().ravel()
        order = np.argsort(-scores, kind='stable')
        order = order[scores[order] > 0][:10]
        for rank, idx in enumerate(order, 1):
            row = docs.iloc[idx]
            print(f'{rank}. {row.title}\n   arXiv {row.arxiv_id}; score={scores[idx]:.4f}\n')
        return
    if args.refresh_presentation:
        summary = json.loads((root / 'results/run_summary.json').read_text(encoding='utf-8'))
        summary['evaluation'] = evaluate_labels(root, summary)
        write_json(root / 'results/run_summary.json', summary)
    else:
        if args.input is None or not args.input.is_file(): parser.error('Укажите существующий файл --input metadata.jsonl')
        summary = run_analysis(args.input, root)
    make_materials(root, summary)




if __name__ == '__main__':
    main()